import os, sys, uuid, time
from pathlib import Path
from datetime import datetime, timedelta
from io import BytesIO
from concurrent.futures import ThreadPoolExecutor
import pytest

os.environ['DATABASE_URL']='sqlite:///'+str(Path(os.getenv('TEMP','.'))/('ponto-test-'+uuid.uuid4().hex+'.db'))
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from fastapi.testclient import TestClient
from sqlalchemy import select
from openpyxl import load_workbook
from backend.main import app
from backend import main, db, rules
from backend.db import Person, Day, Schedule, Setting, Item, Audit, transaction
from backend.security import password_hash

NOW=datetime.fromisoformat('2026-09-03T07:00:00-03:00')
@pytest.fixture(autouse=True)
def clock(monkeypatch):
    monkeypatch.setattr(main,'now',lambda:NOW)
    monkeypatch.setattr(main,'today',lambda:NOW.date().isoformat())
    monkeypatch.setattr(db,'now',lambda:NOW)
    monkeypatch.setattr(db,'today',lambda:NOW.date().isoformat())
    monkeypatch.setattr(rules,'today',lambda:NOW.date().isoformat())
@pytest.fixture
def client():
    with TestClient(app) as c:
        with transaction() as session:
            for table in reversed(db.Base.metadata.sorted_tables):
                if table.name!='settings': session.execute(table.delete())
            session.execute(Setting.__table__.delete().where(Setting.key!='mutex'))
            session.add(Setting(key='geo',data={'lat':-23.67637077,'lon':-46.76243126,'verified':True,'accuracy':100}))
            session.add(Person(id=1,name='Isac',login='suporte',password=password_hash('definitiva1'),role='Suporte',temporary=False,active=True,hired='2026-01-01',details={},session={}))
        c.headers['X-Ponto']='1'
        assert c.post('/api/login',json={'login':'suporte','password':'definitiva1'}).status_code==200
        yield c
def add_person(role='Professor',login='professor'):
    with transaction() as session:
        p=Person(name='Pessoa '+login,login=login,password=password_hash('definitiva1'),role=role,temporary=False,active=True,hired='2026-01-01',details={},session={})
        session.add(p);session.flush();return p.id
def set_schedule(pid,periods,day='2026-01-01'):
    with transaction() as s:s.add(Schedule(person_id=pid,effective=day,specific=False,days={str(i):periods for i in range(7)}))
def punch(c,**kwargs):return c.post('/api/punch',json={'key':uuid.uuid4().hex,'lat':-23.67637077,'lon':-46.76243126,'accuracy':10,**kwargs})
def at(monkeypatch,value):
    t=datetime.fromisoformat(value+'-03:00')
    monkeypatch.setattr(main,'now',lambda:t);monkeypatch.setattr(main,'today',lambda:t.date().isoformat());monkeypatch.setattr(db,'now',lambda:t)

def test_login_first_password_and_masked_profile(client):
    pid=add_person();
    with transaction() as s:
        p=s.get(Person,pid);p.temporary=True;p.password=password_hash('102030')
    with TestClient(app) as c:
        c.headers['X-Ponto']='1'
        assert c.post('/api/login',json={'login':'professor','password':'102030'}).json()['user']['temporary']
        assert c.get('/api/punch/today').status_code==200
        assert c.post('/api/password',json={'current':'102030','password':'segura123'}).status_code==200
        assert c.get('/api/me').json()['user']['login']=='professor'
        assert c.post('/api/profile/login',json={'password':'errada'}).status_code==403
        assert c.get('/api/people').status_code==403
        assert c.get('/api/support/occurrences').status_code==403

def test_tolerance_early_exit_and_idempotency(client,monkeypatch):
    set_schedule(1,[['07:00','15:00']])
    at(monkeypatch,'2026-09-03T06:54:59');assert punch(client).status_code==400
    at(monkeypatch,'2026-09-03T07:06:00');key=uuid.uuid4().hex
    assert punch(client,key=key).status_code==200
    assert punch(client,key=key).status_code==200
    assert len(client.get('/api/punch/today').json()['punches'])==1
    assert punch(client).status_code==409
    at(monkeypatch,'2026-09-03T14:40:00');punch(client)
    r=client.get('/api/punch/today').json()
    assert (r['late'],r['negative'],r['balance'],r['worked'])==(1,21,-21,454)

def test_all_periods_overtime_midnight(client,monkeypatch):
    set_schedule(1,[['05:00','10:00'],['11:00','15:00'],['17:00','20:00']])
    for t in ['05:00','10:00','11:00','15:06','17:00','20:05']:
        at(monkeypatch,'2026-09-03T'+t+':00');assert punch(client).status_code==200
    r=client.get('/api/punch/today').json();assert r['extra']==1 and r['expected']==6
    at(monkeypatch,'2026-09-03T20:10:00');assert punch(client).json()['question']=='overtime'
    assert punch(client,overtime=True).status_code==200
    at(monkeypatch,'2026-09-03T20:40:00');assert punch(client,overtime=True).status_code==200
    assert client.get('/api/punch/today').json()['extra']==31
    at(monkeypatch,'2026-09-04T05:00:00');assert punch(client).status_code==200
    assert len(client.get('/api/punch/today').json()['punches'])==1

def test_forgot_and_absence_arrival(client,monkeypatch):
    set_schedule(1,[['07:00','15:00']])
    assert client.post('/api/absence',json={'person_id':1,'date':'2026-09-03','absent':True}).status_code==200
    at(monkeypatch,'2026-09-03T15:00:00')
    assert punch(client).json()['question']=='forgot'
    r=punch(client,forgot='yes');assert r.status_code==200 and 'falta registrada' in r.json()['message']
    occurrences=client.get('/api/support/occurrences').json()
    assert {x['data']['title'] for x in occurrences}>={'Batida possivelmente esquecida','Falta registrada com comparecimento'}
    assert client.post('/api/absence',json={'person_id':1,'date':'2026-09-03','absent':False}).status_code==200

def test_geofence_and_permissions(client):
    pid=add_person();set_schedule(pid,[['07:00','15:00']])
    with TestClient(app) as c:
        c.headers['X-Ponto']='1';c.post('/api/login',json={'login':'professor','password':'definitiva1'})
        assert c.post('/api/punch',json={'key':uuid.uuid4().hex}).status_code==422
        assert punch(c,lat=-23.67637077,lon=-46.76243126,accuracy=501).status_code==422
        assert punch(c,lat=-23.68,lon=-46.77,accuracy=10).status_code==422
        assert punch(c,lat=-23.67637077,lon=-46.76243126,accuracy=10).status_code==200
        assert c.post('/api/absence',json={'person_id':pid,'date':'2026-09-03','absent':True}).status_code==403
        assert c.get('/api/tickets').json()==[]

def test_schedule_history_immutable(client,monkeypatch):
    set_schedule(1,[['07:00','15:00']]);punch(client)
    r=client.post('/api/people/1/schedules',json={'effective':'2026-09-03','days':{'3':[['08:00','16:00']],'4':[['08:00','16:00']]}})
    assert r.json()['effective']=='2026-09-04'
    assert client.get('/api/punch/today').json()['periods']==[['07:00','15:00']]
    at(monkeypatch,'2026-09-04T08:00:00');assert client.get('/api/punch/today').json()['periods']==[['08:00','16:00']]
    assert client.post('/api/people/1/schedules',json={'effective':'2026-08-03','days':{}}).status_code==400

def test_download_excel_without_closing_and_holiday(client):
    pid=add_person();set_schedule(pid,[['07:00','12:00'],['13:00','17:00'],['18:00','20:00']])
    assert client.post('/api/calendar',json={'date':'2026-08-04','name':'Feriado'}).status_code==200
    assert client.post('/api/absence',json={'person_id':pid,'date':'2026-08-03','absent':True}).status_code==200
    assert client.post('/api/corrections',json={'person_id':pid,'date':'2026-08-05','times':['07:00'],'reason':'Batida esquecida'}).status_code==200
    r=client.post('/api/month/2026-08/export',json={'password':'definitiva1'});assert r.status_code==200
    first=r.content;wb=load_workbook(BytesIO(first));assert len(wb.sheetnames)==4
    sheet=wb['Pontos - Geral'];assert sheet.auto_filter.ref and sheet.freeze_panes=='E2'
    assert any(c.value=='FALTA REGISTRADA' for row in sheet for c in row)
    assert any(c.value=='Feriado' for row in sheet for c in row)
    assert sheet.merged_cells.ranges
    with transaction() as db:
        db.add(Setting(key='month:2026-08',data={'closed':True,'version':1}))
        db.add(Setting(key='export:2026-08:1',data={'xlsx':'obsolete'}))
    assert client.post('/api/calendar',json={'date':'2026-08-04','name':''}).status_code==200
    assert client.post('/api/corrections',json={'person_id':pid,'date':'2026-08-05','times':['07:00','12:00'],'reason':'Ajuste'}).status_code==200
    updated=client.post('/api/month/2026-08/export')
    assert updated.status_code==200
    workbook=load_workbook(BytesIO(updated.content))
    assert not any(c.value=='Feriado' for row in workbook['Pontos - Geral'] for c in row)
    assert 'Fechado' not in workbook.properties.description
    assert client.post('/api/month/2026-08/reopen',json={}).status_code in (404,405)

def test_tickets_delete_and_requests_audit(client):
    pid=add_person()
    with TestClient(app) as c:
        c.headers['X-Ponto']='1';c.post('/api/login',json={'login':'professor','password':'definitiva1'})
        assert c.post('/api/tickets',json={'message':'Preciso de ajuda'}).status_code==200
        assert client.get('/api/tickets/unread').json()=={'count':1,'names':['Pessoa professor']}
        ticket=c.get('/api/tickets').json()[0]['id']
        assert c.delete('/api/tickets/'+str(ticket)).status_code==403
        assert client.post(f'/api/tickets/{ticket}/message',json={'message':'Vou verificar'}).status_code==200
        assert client.get('/api/tickets/unread').json()['count']==1
        assert client.post('/api/tickets/read',json={}).json()=={'read':1}
        assert client.get('/api/tickets/unread').json()=={'count':0,'names':[]}
        assert len(c.get('/api/tickets').json()[0]['data']['messages'])==2
        assert client.delete('/api/tickets/'+str(ticket)).status_code==200
        assert c.get('/api/tickets').json()==[]
        c.post('/api/requests',json={'type':'profile','field':'name','value':'Nome corrigido','reason':'Acentuação'})
        req=c.get('/api/requests').json()[0]['id']
        assert c.get('/api/me').json()['user']['name']!='Nome corrigido'
        assert client.post(f'/api/requests/{req}/decide',json={'approve':True,'reason':'Conferido'}).status_code==200
        assert c.get('/api/me').json()['user']['name']=='Nome corrigido'
    assert any(r['action']=='Aprovar alteração cadastral' for r in client.get('/api/audit?month=2026-09').json())

def test_single_session_immediate_and_same_browser(client):
    before=client.cookies.get('ponto_session')
    assert client.post('/api/login',json={'login':'suporte','password':'definitiva1'}).status_code==200
    assert client.cookies.get('ponto_session')==before
    assert client.get('/api/me').status_code==200
    with TestClient(app) as c:
        c.headers['X-Ponto']='1'
        r=c.post('/api/login',json={'login':'suporte','password':'definitiva1'})
        assert r.status_code==200 and 'pending' not in r.json()
        assert client.get('/api/me').status_code==401
        assert c.get('/api/me').status_code==200


def test_double_requests_and_inactivation(client):
    set_schedule(1,[['07:00','15:00']]);key=uuid.uuid4().hex
    with ThreadPoolExecutor(max_workers=2) as pool:
        results=list(pool.map(lambda _:punch(client,key=key),range(2)))
    assert all(r.status_code==200 for r in results)
    assert len(client.get('/api/punch/today').json()['punches'])==1
    pid=add_person()
    assert client.post(f'/api/people/{pid}/deactivate',json={'date':'2026-09-03'}).status_code==200
    assert client.post('/api/login',json={'login':'professor','password':'definitiva1'}).status_code==401
    with transaction() as s: assert s.get(Person,pid) is not None


def test_pool_configuration():
    assert db.connection_args('postgresql+psycopg://localhost/test')['prepare_threshold'] is None
    assert 'prepare_threshold' not in db.connection_args('sqlite:///test.db')

def test_late_first_and_four_punches(client,monkeypatch):
    set_schedule(1,[['07:00','12:00'],['13:00','17:00']])
    for value in ['09:00','12:00','13:00','17:00']:
        at(monkeypatch,'2026-09-03T'+value+':00')
        result=punch(client)
        assert result.status_code==200 and 'question' not in result.json()
    result=client.get('/api/punch/today').json()
    assert len(result['punches'])==4 and result['late']==115

def test_geo_confirmation_persists(client):
    settings=client.get('/api/support/settings').json()
    settings.update(verified=False,password='definitiva1')
    assert client.put('/api/support/settings',json=settings).status_code==200
    set_schedule(1,[['07:00','15:00']])
    assert punch(client,test=False,lat=settings['lat'],lon=settings['lon'],accuracy=10).status_code==409
    settings['verified']=True
    assert client.put('/api/support/settings',json=settings).status_code==200
    db.init()
    assert client.get('/api/support/settings').json()['verified'] is True
    assert punch(client,test=False,lat=settings['lat'],lon=settings['lon'],accuracy=10).status_code==200

def test_top_profile_and_support_no_self_ticket(client):
    assert client.post('/api/tickets',json={'message':'teste'}).status_code==403
    assert client.put('/api/profile',json={'phone':'11 98564-1624'}).status_code==200
    assert client.get('/api/me').json()['user']['details']['phone']=='11985641624'
    add_person(role='Diretoria',login='diretoria')
    with TestClient(app) as c:
        c.headers['X-Ponto']='1'
        c.post('/api/login',json={'login':'diretoria','password':'definitiva1'})
        assert c.get('/api/support/settings').status_code==200
        assert c.get('/api/tickets?management=true').status_code==403
        assert c.put('/api/profile',json={'name':'Diretora Ana'}).status_code==200
    add_person(role='Administração',login='admin')
    with TestClient(app) as c:
        c.headers['X-Ponto']='1'
        c.post('/api/login',json={'login':'admin','password':'definitiva1'})
        assert c.get('/api/support/settings').status_code==403
        assert c.put('/api/profile',json={'name':'Teste'}).status_code==403

def test_health_head(client):
    assert client.head('/health').status_code==200


def test_requested_support_four_times(client,monkeypatch):
    periods=[['07:30','11:00'],['12:00','17:30']]
    result=client.post('/api/people/1/schedules',json={'effective':'2026-09-03','days':{'3':periods}})
    assert result.json()['deferred'] is False
    assert client.get('/api/punch/today').json()['expected']==4
    for value in ['07:30','11:00','12:00','17:30']:
        at(monkeypatch,'2026-09-03T'+value+':00')
        response=punch(client)
        assert response.status_code==200 and 'question' not in response.json()
    assert client.get('/api/punch/today').json()['status']=='Completo'
    result=client.post('/api/people/1/schedules',json={'effective':'2026-09-03','days':{'4':periods}})
    assert result.json()=={'effective':'2026-09-04','deferred':True}

def test_schedule_query_reused_within_request(client):
    from sqlalchemy import event
    pid=add_person();set_schedule(pid,[['07:30','11:00'],['12:00','17:30']])
    queries=[]
    def count(conn,cursor,statement,parameters,context,executemany):
        if 'FROM schedules' in statement: queries.append(statement)
    event.listen(db.engine,'before_cursor_execute',count)
    try:
        result=client.get(f'/api/records?start=2026-08-01&end=2026-08-31&person_id={pid}')
        assert result.status_code==200
        assert len(queries)==1
    finally:event.remove(db.engine,'before_cursor_execute',count)


def test_global_qr_uses_authenticated_person(client,monkeypatch):
    pid=add_person(login='qruser');set_schedule(pid,[['07:00','12:00'],['13:00','17:00']])
    generated=client.post('/api/system/qr',json={})
    assert generated.status_code==200
    token=generated.json()['url'].split('/q/')[1]
    with TestClient(app) as anonymous:
        anonymous.headers['X-Ponto']='1'
        assert anonymous.get('/api/qr/'+token).status_code==401
        assert anonymous.post('/api/login',json={'login':'qruser','password':'definitiva1'}).status_code==200
        state=anonymous.get('/api/qr/'+token)
        assert state.status_code==200 and state.json()['next']=='Entrada'
        result=anonymous.post('/api/qr/'+token+'/punch',json={'key':uuid.uuid4().hex,'lat':-23.67637077,'lon':-46.76243126,'accuracy':10})
        assert result.status_code==200 and result.json()['confirmation']=='Pessoa, sua entrada foi registrada com sucesso'
    replaced=client.post('/api/system/qr',json={}).json()['url'].split('/q/')[1]
    assert client.get('/api/qr/'+token).status_code==404
    assert client.get('/api/qr/'+replaced).status_code==200
    assert client.get('/api/system/qr/image').headers['content-type']=='image/png'

def test_admission_and_password_permission(client):
    pid=add_person(login='labels')
    person=next(x for x in client.get('/api/people').json() if x['id']==pid)
    person['hired']='1999-02-01'
    assert client.put(f'/api/people/{pid}',json=person).status_code==200
    assert next(x for x in client.get('/api/people').json() if x['id']==pid)['hired']=='1999-02-01'
    admin=add_person(role='Administração',login='middle')
    with TestClient(app) as c:
        c.headers['X-Ponto']='1';c.post('/api/login',json={'login':'middle','password':'definitiva1'})
        assert c.post(f'/api/people/{pid}/reset',json={}).status_code==403
        assert c.get('/api/support/occurrences').status_code==403
        assert c.get('/api/tickets?management=true').status_code==403



def test_optional_temporary_password_and_four_characters(client):
    pid=add_person(login='temporary')
    with transaction() as db:
        p=db.get(Person,pid);p.temporary=True;p.password=password_hash('102030')
    with TestClient(app) as c:
        c.headers['X-Ponto']='1'
        assert c.post('/api/login',json={'login':'temporary','password':'102030'}).status_code==200
        assert c.get('/api/me').status_code==200
        assert c.get('/api/punch/today').status_code==200
        assert c.post('/api/password',json={'password':'abc'}).status_code==400
        assert c.post('/api/password',json={'password':'aB12'}).status_code==200
        assert c.get('/api/me').json()['user']['temporary'] is False

def test_recovery_email_single_use_expiry_and_cooldown(client,monkeypatch):
    pid=add_person(login='recoverme')
    with transaction() as db:
        p=db.get(Person,pid);p.details={'email':'pessoa@example.test'}
    sent=[]
    class SMTP:
        def __init__(self,*args,**kwargs):pass
        def __enter__(self):return self
        def __exit__(self,*args):pass
        def starttls(self,**kwargs):pass
        def login(self,*args):pass
        def send_message(self,message):sent.append(message)
    monkeypatch.setenv('SMTP_HOST','smtp.example.test');monkeypatch.setenv('SMTP_FROM','sistema@example.test');monkeypatch.setattr(main.smtplib,'SMTP',SMTP)
    neutral='Se o login estiver cadastrado e possuir e-mail, uma credencial temporária será enviada.'
    assert client.post('/api/recover',json={'login':'recoverme'}).json()['message']==neutral
    credential=sent[0].get_content().splitlines()[0]
    assert len(credential)==8 and credential.isalnum()
    with transaction() as db:
        stored=db.get(Setting,f'recovery:{pid}').data
        assert credential not in str(stored) and stored['hash']==main.digest(credential)
    assert client.post('/api/recover',json={'login':'recoverme'}).json()['message']==neutral and len(sent)==1
    with TestClient(app) as c:
        c.headers['X-Ponto']='1';response=c.post('/api/login',json={'login':'recoverme','password':credential})
        assert response.status_code==200 and response.json()['user']['temporary']
    with TestClient(app) as c:
        c.headers['X-Ponto']='1';assert c.post('/api/login',json={'login':'recoverme','password':credential}).status_code==401
    with transaction() as db:
        rate=db.get(Setting,'recover-rate:'+main.digest('recoverme'));rate.data={'until':0}
    client.post('/api/recover',json={'login':'recoverme'});expired=sent[-1].get_content().splitlines()[0]
    with transaction() as db:
        row=db.get(Setting,f'recovery:{pid}');row.data={**row.data,'expires':0}
    with TestClient(app) as c:
        c.headers['X-Ponto']='1';assert c.post('/api/login',json={'login':'recoverme','password':expired}).status_code==401

def test_support_location_exception_is_per_account_and_audited(client):
    assert client.get('/api/support/my-location').json()=={'required':True}
    assert client.put('/api/support/my-location',json={'required':False}).json()=={'required':False}
    set_schedule(1,[['07:00','15:00']]);assert client.post('/api/punch',json={'key':uuid.uuid4().hex}).status_code==200
    other=add_person(role='Suporte',login='othersupport');set_schedule(other,[['07:00','15:00']])
    with TestClient(app) as c:
        c.headers['X-Ponto']='1';c.post('/api/login',json={'login':'othersupport','password':'definitiva1'})
        assert c.get('/api/support/my-location').json()=={'required':True}
        assert c.post('/api/punch',json={'key':uuid.uuid4().hex}).status_code==422
    assert any(a['action']=='Alterar exigência de localização própria' for a in client.get('/api/audit?month=2026-09').json())

def test_grade_times_are_manual_and_do_not_change_schedule(client,monkeypatch):
    pid=add_person(login='teacher');set_schedule(pid,[['07:10','08:50'],['09:10','12:30']])
    subjects={'1º Ano A':['Matemática'],'2º Ano A':['Matemática'],'3º Ano A':['Artes'],'4º Ano A':['Artes'],'5º Ano A':['Português']}
    result=client.put(f'/api/people/{pid}/subjects',json={'subjects':subjects})
    assert result.status_code==200 and result.json()['subjects']==subjects
    lessons=[
        {'day':'0','start':'07:10','end':'08:00','class':'1º Ano A','subject':'Matemática'},
        {'day':'0','start':'08:00','end':'08:50','class':'2º Ano A','subject':'Matemática'},
        {'day':'0','start':'09:10','end':'10:00','class':'3º Ano A','subject':'Artes'},
        {'day':'0','start':'10:00','end':'10:50','class':'4º Ano A','subject':'Artes'},
        {'day':'0','start':'10:50','end':'11:40','class':'5º Ano A','subject':'Português'},
    ]
    assert client.put(f'/api/people/{pid}/lessons',json={'lessons':lessons}).status_code==200
    invalid={**lessons[0],'end':'07:00'}
    assert client.put(f'/api/people/{pid}/lessons',json={'lessons':[invalid]}).status_code==400
    at(monkeypatch,'2026-09-07T07:10:00')
    with TestClient(app) as teacher:
        teacher.headers['X-Ponto']='1';teacher.post('/api/login',json={'login':'teacher','password':'definitiva1'})
        point=teacher.get('/api/punch/today').json()
        assert point['periods']==[['07:10','08:50'],['09:10','12:30']] and point['expected']==4
    storage=client.get('/api/system/storage')
    assert storage.status_code==200 and storage.json()['used_bytes']>0 and storage.json()['limit_bytes']>0

def test_one_login_can_choose_professor_or_administrator(client,monkeypatch):
    pid=add_person(login='edson');set_schedule(pid,[['07:10','08:50']])
    with transaction() as db:
        person=db.get(Person,pid);person.details={**person.details,'roles':['Professor','Administração']}
    at(monkeypatch,'2026-09-03T07:10:00')
    with TestClient(app) as user:
        user.headers['X-Ponto']='1'
        login=user.post('/api/login',json={'login':'edson','password':'definitiva1'}).json()
        assert login['user']['access_required'] and login['punch'] is None
        assert user.get('/api/people').status_code==409
        chosen=user.post('/api/access',json={'role':'Professor'}).json()
        assert chosen['user']['role']=='Professor' and not chosen['user']['access_required']
        assert user.get('/api/people').status_code==403
        assert punch(user).status_code==200
        user.post('/api/logout',json={})
        user.post('/api/login',json={'login':'edson','password':'definitiva1'})
        assert user.post('/api/access',json={'role':'Administração'}).json()['user']['role']=='Administração'
        assert user.get('/api/people').status_code==200
    workbook=load_workbook(BytesIO(client.post('/api/month/2026-09/export',json={'password':'definitiva1'}).content))
    names=[cell.value for cell in workbook['Pontos - Geral']['C']]
    assert 'Pessoa edson (Professor)' in names


def test_teacher_export_and_notebook_location(client):
    pid=add_person(login='exportteacher')
    with transaction() as session:
        person=session.get(Person,pid)
        person.details={'subjects_by_class':{'1º Ano A':['Português','Matemática'],'2º Ano A':['Arte']}}
    response=client.get('/api/reports/teachers')
    assert response.status_code==200
    wb=load_workbook(BytesIO(response.content)); sheet=wb['Professores']
    assert list(sheet.values)==[('Professor','Turma','Matéria'),('Pessoa exportteacher','1º Ano A','Português'),('Pessoa exportteacher','1º Ano A','Matemática'),('Pessoa exportteacher','2º Ano A','Arte')]
    assert sheet.auto_filter.ref=='A1:C4'
    set_schedule(pid,[['07:00','15:00']])
    with TestClient(app) as user:
        user.headers['X-Ponto']='1';user.post('/api/login',json={'login':'exportteacher','password':'definitiva1'})
        assert user.get('/api/reports/teachers').status_code==403
        assert punch(user,lat=-23.68,lon=-46.77,accuracy=400).status_code==422
        assert punch(user,lat=-23.67637077,lon=-46.76243126,accuracy=400).status_code==200


def test_grade_saves_new_subjects_atomically(client):
    pid=add_person(login='newgrade')
    lesson={'day':'0','start':'07:10','end':'08:00','class':'1º Ano A','subject':'Português'}
    # A matéria aparece na tela antes de ser salva separadamente.
    response=client.put(f'/api/people/{pid}/lessons',json={'lessons':[lesson]})
    assert response.status_code==400 and 'Segunda-feira, 1ª aula' in response.json()['detail']
    subjects={'1º Ano A':['Português']}
    assert client.put(f'/api/people/{pid}/lessons',json={'lessons':[lesson],'subjects':subjects}).status_code==200
    with transaction() as db:
        saved=db.get(Person,pid).details
        assert saved['lessons']==[lesson] and saved['subjects_by_class']==subjects
    invalid={**lesson,'day':'1','end':'06:00'}
    response=client.put(f'/api/people/{pid}/lessons',json={'lessons':[lesson,invalid],'subjects':{'1º Ano A':['Português','Arte']}})
    assert response.status_code==400 and 'Terça-feira, 1ª aula' in response.json()['detail']
    with transaction() as db:
        saved=db.get(Person,pid).details
        assert saved['lessons']==[lesson] and saved['subjects_by_class']==subjects


def test_people_indicates_registered_schedule(client):
    missing=add_person(login='missing');empty=add_person(login='empty');registered=add_person(login='registered')
    set_schedule(empty,[]);set_schedule(registered,[['07:10','12:30']],day='2026-10-01')
    rows={row['id']:row for row in client.get('/api/people').json()}
    assert rows[missing]['has_schedule'] is False
    assert rows[empty]['has_schedule'] is False
    assert rows[registered]['has_schedule'] is True


def test_database_reset_requires_support_and_confirmation(client):
    pid=add_person(login='resetteacher');set_schedule(pid,[['07:00','12:00']])
    assert client.post('/api/absence',json={'person_id':pid,'date':'2026-08-03','absent':True}).status_code==200
    with TestClient(app) as teacher:
        teacher.headers['X-Ponto']='1';teacher.post('/api/login',json={'login':'resetteacher','password':'definitiva1'})
        teacher.post('/api/tickets',json={'message':'Teste'})
        assert teacher.post('/api/system/reset',json={'password':'definitiva1','confirmation':'APAGAR'}).status_code==403
        assert client.post('/api/system/reset',json={'password':'errada','confirmation':'APAGAR'}).status_code==403
        assert client.post('/api/system/reset',json={'password':'definitiva1','confirmation':''}).status_code==400
        with transaction() as db: assert db.get(Person,pid) is not None
        assert client.post('/api/system/reset',json={'password':'definitiva1','confirmation':'APAGAR'}).status_code==200
        assert teacher.get('/api/me').status_code==401
    assert client.get('/api/me').status_code==401
    with transaction() as db:
        assert list(db.scalars(select(Person.id)))==[1]
        for model in (Day,Schedule,Item): assert db.scalar(select(model).limit(1)) is None
        assert list(db.scalars(select(Setting.key)))==['mutex']
        assert list(db.scalars(select(Audit.action)))==['Apagar dados do sistema']
    assert client.post('/api/login',json={'login':'suporte','password':'definitiva1'}).status_code==200


def test_reset_clears_legacy_and_rolls_back_on_failure(client,monkeypatch):
    from sqlalchemy import text
    pid=add_person(login='preserved_on_failure')
    with transaction() as session:
        session.execute(text('CREATE TABLE usuarios (id INTEGER PRIMARY KEY, nome TEXT)'))
        session.execute(text('CREATE TABLE registros_ponto (id INTEGER PRIMARY KEY, usuario_id INTEGER REFERENCES usuarios(id))'))
        session.execute(text('CREATE TABLE unrelated (id INTEGER PRIMARY KEY)'))
        session.execute(text("INSERT INTO usuarios VALUES (1, 'Antigo')"))
        session.execute(text('INSERT INTO registros_ponto VALUES (1, 1)'))
        session.execute(text('INSERT INTO unrelated VALUES (1)'))
    original=main.audit
    def reject(*args,**kwargs): raise RuntimeError('Falha simulada antes de confirmar a limpeza')
    monkeypatch.setattr(main,'audit',reject)
    response=client.post('/api/system/reset',json={'password':'definitiva1','confirmation':'APAGAR'})
    assert response.status_code==500
    with transaction() as session:
        assert session.get(Person,pid) is not None
        assert session.scalar(text('SELECT count(*) FROM usuarios'))==1
        assert session.scalar(text('SELECT count(*) FROM registros_ponto'))==1
    monkeypatch.setattr(main,'audit',original)
    assert client.post('/api/system/reset',json={'password':'definitiva1','confirmation':'APAGAR'}).status_code==200
    with transaction() as session:
        assert session.get(Person,pid) is None
        assert session.scalar(text('SELECT count(*) FROM usuarios'))==0
        assert session.scalar(text('SELECT count(*) FROM registros_ponto'))==0
        assert session.scalar(text('SELECT count(*) FROM unrelated'))==1
        session.execute(text('DROP TABLE registros_ponto'))
        session.execute(text('DROP TABLE usuarios'))
        session.execute(text('DROP TABLE unrelated'))
    assert client.post('/api/login',json={'login':'suporte','password':'definitiva1'}).status_code==200
